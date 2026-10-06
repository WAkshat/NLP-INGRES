"""Builds the mid-term PPT on a copy of the university template (cover slide, master band and logo kept)."""
import copy
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

R = Path(__file__).resolve().parents[2]
TPL = Path(r"C:\Users\PC\Downloads\Mid Term PPT (1).pptx")
OUT = R / "reports/midterm/Midterm_PPT_INGRES_Text2SQL.pptx"
FIG = R / "reports/midterm/fig"
TITLE = "INGRES-Text2SQL: A Multilingual Natural-Language Interface to India's Groundwater Assessment Data"
REPO = "github.com/WAkshat/NLP-INGRES"
NAVY, ORANGE, RED = RGBColor(0x2E, 0x3A, 0x59), RGBColor(0xE0, 0x70, 0x3A), RGBColor(0x8B, 0x2A, 0x2A)
L, T, W, H = Emu(838200), Emu(1690000), Emu(10515600), Emu(4150000)   # content area: below title, above bottom-right logos (~5.9M)
FONT = "Calibri"

prs = Presentation(TPL)
S = list(prs.slides)
A = json.loads((R / "experiments/baselines/A_keyword_template/metrics.json").read_text(encoding="utf-8"))
SL = json.loads((R / "experiments/schema_linking/baselines.json").read_text(encoding="utf-8"))["results"]


def set_title(slide, text, size=34):
    t = slide.shapes.title
    t.text = text
    for p in t.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), True, NAVY, FONT


def bullets(slide, items, left=L, top=T, width=W, height=H, size=18):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        sub = it.startswith("-")
        text = it.lstrip("- ")
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(342900 * (2 if sub else 1)))
        pPr.set("indent", "-228600")
        bu = pPr.makeelement(qn("a:buChar"), {"char": "\u2013" if sub else "\u2022"})
        pPr.append(bu)
        p.space_after = Pt(6)
        for k, seg in enumerate(text.split("**")):
            if seg:
                r = p.add_run()
                r.text = seg
                r.font.size = Pt(size - 2 if sub else size)
                r.font.bold = k % 2 == 1
                r.font.name = FONT
                if k % 2 == 1:
                    r.font.color.rgb = NAVY
    return tb


def table(slide, header, rows, widths, top=T, size=12, height=None):
    tbl = slide.shapes.add_table(len(rows) + 1, len(header), L, top, W, height or Emu(380000 * (len(rows) + 1))).table
    tot = sum(widths)
    for j, w in enumerate(widths):
        tbl.columns[j].width = Emu(int(W * w / tot))
    for j, h in enumerate(header):
        c = tbl.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        for p in c.text_frame.paragraphs:
            for r in p.runs:
                r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), True, RGBColor(255, 255, 255), FONT
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = tbl.cell(i, j)
            c.text = str(v)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xF3, 0xF4, 0xF7) if i % 2 else RGBColor(255, 255, 255)
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
                    r.font.name = FONT
                    r.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    return tbl


def picture(slide, path, top=T, max_w=W, max_h=H):
    pic = slide.shapes.add_picture(str(path), L, top)
    ratio = min(max_w / pic.width, max_h / pic.height)
    pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
    pic.left = int(L + (W - pic.width) / 2)
    return pic


def new_slide(title):
    s = prs.slides.add_slide(prs.slide_layouts[5])   # "Title Only", same as template content slides
    set_title(s, title)
    return s


def note(slide, text, top=Emu(5900000)):
    tb = slide.shapes.add_textbox(L, top, Emu(8600000), Emu(350000))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text, r.font.size, r.font.italic, r.font.color.rgb, r.font.name = text, Pt(12), True, RGBColor(0x60, 0x60, 0x60), FONT


# -------- slide 2: title & team (template shapes)
s2 = S[1]
for sh in s2.shapes:
    if sh.name == "object 2":
        set_title(s2, TITLE, 30)
        sh.width = Emu(9300000)   # keep clear of the logo at top-right
    elif sh.name == "object 3":
        sh.text_frame.text = "Team Member(s):\n<Name (Roll No.)>\n<Name (Roll No.)>"
    elif sh.name == "object 4":
        sh.text_frame.text = "Supervisor(s): <Name, Designation>"
    elif sh.name == "object 5":
        sh.text_frame.text = "Department of Computer Science and Engineering"
    if sh.has_text_frame:
        for p_ in sh.text_frame.paragraphs:
            for r_ in p_.runs:
                r_.font.name = FONT

# -------- slide 3: introduction & motivation
set_title(S[2], "Introduction & Motivation")
bullets(S[2], [
    "Groundwater supplies most of India's irrigation; CGWB assesses ~7,000 blocks / mandals / taluks every cycle "
    "(recharge, extraction, **stage of extraction**, Safe to Over-Exploited category)",
    "Results are published through **INGRES** (CGWB + IIT Hyderabad): a GIS portal with maps and tables, English only",
    "Real users ask questions like **\"Ludhiana ke kitne blocks over-exploited hain?\"** or "
    "\"Which Punjab blocks moved from Critical to Over-Exploited since 2022?\"",
    "**Text-to-SQL** can answer these, but research is English-centric: no benchmark or system covers Indian scripts, "
    "Hinglish or ambiguous Indian place names",
    "This project: a **multilingual (English, Hindi, Hinglish, Tamil) Text-to-SQL system over official INGRES data**, "
    "studied as research with a benchmark, baselines and ablations",
])

# -------- slide 7 (template): problem statement
for sh in list(S[6].shapes):
    if sh.name == "object 3":                  # template's example link
        sh._element.getparent().remove(sh._element)
set_title(S[6], "Problem Statement")
bullets(S[6], [
    "Design, build and evaluate a natural-language interface that converts questions about India's groundwater "
    "assessments, written in **English, Hindi, romanised Hinglish or Tamil**, into **correct and safe SQL** over the "
    "official INGRES data, and returns answers whose **numbers are grounded in the query result**.",
    "Measure, by **execution accuracy** on a purpose-built benchmark, how much each component contributes:",
    "- schema linking across scripts   - geographic entity resolution   - tokenizer adaptation",
    "- numeric grounding verification   - agentic decomposition for multi-step questions",
    "Handle **code-mixed and noisy input** and **historical comparisons across assessment cycles** with methodology caveats",
], size=20)

# -------- slide 8 (template): objectives
set_title(S[7], "Objectives")
bullets(S[7], [
    "**O1** Validated, reproducible database of official INGRES assessments for 5 cycles  \u2714 done",
    "**O2** INGRES-Bench: multilingual Text-to-SQL benchmark with leakage-free splits and a hard test set  \u2714 done",
    "**O3** Baselines: keyword/template, BM25, multilingual embeddings, frontier LLM  (in progress)",
    "**O4** Learned multilingual schema linker with hard negatives (recall@k, MRR per language)",
    "**O5** Hierarchical geographic entity resolver for misspelled, transliterated and ambiguous names",
    "**O6** Tokenizer fertility and code-mixing study: link to accuracy; compare translation vs direct vs vocabulary adaptation",
    "**O7** Numeric grounding verifier: every number in an answer traced to the SQL result (pass / regenerate / abstain)",
    "**O8** Single-shot vs agentic decomposition: accuracy, SQL calls, latency, cost",
], size=16)

# -------- literature survey (template slide 5)
set_title(S[4], "Literature Survey")
table(S[4], ["Work", "Approach", "Limitation for our setting"], [
    ["Spider (Yu et al., EMNLP 2018)", "Cross-domain Text-to-SQL benchmark, 200 databases", "English only"],
    ["BIRD (Li et al., NeurIPS 2023)", "Large benchmark with dirty values, external knowledge", "English only"],
    ["RAT-SQL (ACL 2020), RESDSQL (AAAI 2023)", "Schema-aware encoders; decoupled schema linking", "Trained on English questions"],
    ["DIN-SQL (NeurIPS 2023), DAIL-SQL (PVLDB 2024)", "LLM prompting, decomposition, self-correction", "English only; decomposes every question"],
    ["MAC-SQL (COLING 2025), CHESS (2024)", "Multi-agent / pipeline LLM systems", "Cost; per-question-type benefit not isolated"],
    ["CSpider (EMNLP 2019), MultiSpider (AAAI 2023)", "Translated multilingual Text-to-SQL", "No Indian scripts, no code-mixing"],
    ["Petrov et al. (NeurIPS 2023), Rust et al. (ACL 2021)", "Tokenizer fertility / unfairness across languages", "Not linked to Text-to-SQL accuracy"],
    ["GLUECoS (ACL 2020), CMI (LREC 2016)", "Code-switched NLP benchmark; code-mixing index", "No semantic parsing task"],
    ["FActScore (EMNLP 2023), hallucination survey", "Measuring unsupported facts in generated text", "Not for numbers derived from SQL"],
], [3.2, 3.6, 3.0], size=11)

# -------- existing systems & comparison (template slide 5 duplicate -> new)
s_cmp = new_slide("Existing Systems & Comparison")
table(s_cmp, ["Parameter", "INGRES portal", "English Text-to-SQL", "Multilingual (MultiSpider)", "Proposed"], [
    ["Natural-language questions", "No", "Yes", "Yes", "Yes"],
    ["Indian languages / scripts", "No", "No", "No", "Hindi, Tamil"],
    ["Hinglish (romanised code-mixed)", "No", "No", "No", "Yes"],
    ["Ambiguous place names", "Manual navigation", "Not addressed", "Not addressed", "Gazetteer + context resolver"],
    ["Cross-year comparison with caveats", "Manual", "No", "No", "Crosswalk + caveats"],
    ["Numbers verified against data", "N/A", "Rarely", "No", "Grounding verifier"],
    ["Public benchmark / evaluation", "N/A", "Yes", "Yes", "INGRES-Bench"],
    ["Official government data", "Yes", "No", "No", "Yes"],
], [3.0, 1.7, 1.8, 2.0, 2.4], size=13)
note(s_cmp, "Public INGRES chatbot prototypes we inspected do not document their data source or report any evaluation.")

# -------- research gap (template slide 6)
set_title(S[5], "Research Gap")
bullets(S[5], [
    "**No multilingual benchmark on Indian public data:** no Text-to-SQL benchmark covers Hindi, Hinglish or a Dravidian language",
    "**Cross-script schema linking fails:** our baselines show BM25 recall@3 of ~0.01 for Hindi and Tamil script vs 0.45 for English",
    "**Ambiguous, inconsistently spelled places:** 169 unit names are shared across states ('Ramnagar' x10); "
    "INGRES itself respells units across cycles (NOWGAON to NOWGONG)",
    "**Unsafe cross-year joins:** INGRES re-issues unit IDs, and some states changed the unit type (Tamil Nadu firka to taluk)",
    "**Unverified numbers in LLM answers**, which is unacceptable for policy data",
    "**The benefit of agents is not isolated** by question type; they add cost and latency",
    "How we address them: INGRES-Bench, learned schema linker, entity resolver, crosswalk + caveats, grounding verifier, targeted agent",
], size=16)

# -------- methodology
s_m = new_slide("Proposed Methodology")
picture(s_m, FIG / "architecture.png")

# -------- work completed
s_w = new_slide("Work Completed (Mid-term)")
ex = A["execution_accuracy"]
bullets(s_w, [
    "**Data access:** found INGRES's public JSON endpoint; crawled 5,299 responses with 0 failures",
    "**Database:** 5 cycles (GWRA 2020, 2022-2025); national totals match PIB figures within 0.1% (automated test)",
    "**Crosswalk:** links units across cycles despite re-issued IDs; fuzzy links 37/40 correct on manual check",
    "**INGRES-Bench v1:** 371 items x 4 languages = 1,484 NL-SQL pairs; 0 QC failures on 8 checks",
    f"**Baselines:** keyword/template EX {100 * ex:.1f}%; schema recall@3: BM25 0.267, mE5 0.409 (random 0.028)",
    "**Safety:** read-only SQL executor; 13 injection / write tests pass",
    f"Code: **{REPO}**",
], width=Emu(5200000), size=15)
pic = s_w.shapes.add_picture(str(FIG / "baseline_results.png"), Emu(6150000), T, width=Emu(5300000))

# -------- feasibility & outcomes (template slide 4)
set_title(S[3], "Feasibility & Expected Outcomes")
bullets(S[3], [
    "**Technical:** data access verified; 22 MB SQLite DB; CPU for retrieval, RTX 5060 (8 GB) for fine-tuning",
    "**Tools:** Python, SQLite, pandas, scikit-learn, sentence-transformers, multilingual-e5, Gemini API, Git/GitHub",
    "**Economic:** open-source tools, public data, free-tier LLM with on-disk caching; zero cost",
    "**Operational / schedule:** fully scripted, reproducible pipeline; riskiest phases (data, benchmark) already done",
], height=Emu(2000000), size=15)
bullets(S[3], [
    "**Expected outcomes**",
    "- Validated INGRES database and data-access report (achieved)",
    "- INGRES-Bench: first multilingual, code-mixed Text-to-SQL benchmark on Indian government data (achieved)",
    "- Measured contribution of each component via ablations, including negative results",
    "- Working read-only prototype with SQL, result, grounded answer and cross-year caveats",
], top=Emu(3800000), height=Emu(2000000), width=Emu(9000000), size=15)

# -------- gantt (template slide 9)
set_title(S[8], "Project Timeline (Gantt Chart)")
picture(S[8], FIG / "gantt.png")

# -------- responsibility (template slide 10)
set_title(S[9], "Responsibility Chart")
table(S[9], ["Task", "<Member 1>", "<Member 2>", "<Member 3>", "<Member 4>"], [
    ["Data access, crawling and audit", "R", "A", "C", "I"],
    ["Canonical database and crosswalk", "R", "C", "A", "I"],
    ["INGRES-Bench construction and QC", "A", "R", "R", "C"],
    ["Baselines and evaluation harness", "C", "R", "A", "R"],
    ["Schema linking and tokenization study", "C", "A", "R", "C"],
    ["Entity resolution and numeric grounding", "R", "C", "C", "A"],
    ["Agent, UI and documentation", "A", "C", "C", "R"],
], [4.2, 1.5, 1.5, 1.5, 1.5], size=14)
note(S[9], "R = responsible, A = accountable, C = consulted, I = informed")

# -------- references
s_r = new_slide("References")
bullets(s_r, [
    "Yu et al., Spider, EMNLP 2018  |  Li et al., BIRD, NeurIPS 2023  |  Wang et al., RAT-SQL, ACL 2020",
    "Pourreza & Rafiei, DIN-SQL, NeurIPS 2023  |  Gao et al., DAIL-SQL, PVLDB 17(5) 2024  |  Li et al., RESDSQL, AAAI 2023",
    "Min et al., CSpider, EMNLP-IJCNLP 2019  |  Dou et al., MultiSpider, AAAI 2023",
    "Wang et al., MAC-SQL, COLING 2025  |  Talaei et al., CHESS, arXiv:2405.16755",
    "Petrov et al., Tokenizer unfairness, NeurIPS 2023  |  Rust et al., How good is your tokenizer?, ACL 2021",
    "Khanuja et al., GLUECoS, ACL 2020  |  Gambäck & Das, Code-switching level, LREC 2016",
    "Gala et al., IndicTrans2, TMLR 2023  |  Wang et al., Multilingual E5, arXiv:2402.05672",
    "Robertson & Zaragoza, BM25 and beyond, FnTIR 2009  |  Min et al., FActScore, EMNLP 2023  |  Ji et al., ACM CSUR 2023",
    "CGWB, GEC-2015 report (2017)  |  PIB releases on Dynamic Ground Water Resources 2022-2025  |  INGRES portal, ingres.iith.ac.in",
], size=14)

# -------- reorder slides as in the midterm instructions
order = [S[0], S[1], S[2], S[6], S[7], S[4], s_cmp, S[5], s_m, s_w, S[3], S[8], S[9], s_r, S[10], S[11]]
lst = prs.slides._sldIdLst
ids = {sl.slide_id: el for el in list(lst) for sl in prs.slides if sl.slide_id == int(el.get("id"))}
els = [ids[s.slide_id] for s in order]
for el in list(lst):
    lst.remove(el)
for el in els:
    lst.append(el)
prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
